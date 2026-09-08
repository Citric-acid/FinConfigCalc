from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from fin_config_calc.utils.io_dataframe_excel import read_excel, write_excel


def test_read_excel_allows_empty_cells(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "amount", "note"])
    worksheet.append(["revenue", None, None])
    workbook.save(input_path)

    dataframe = read_excel(input_path)

    assert dataframe.shape == (1, 3)
    assert dataframe.row(0) == ("revenue", None, None)


def test_read_excel_ignores_formatted_trailing_empty_columns(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "amount"])
    worksheet.append(["revenue", 100])
    worksheet["D1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.save(input_path)

    dataframe = read_excel(input_path)

    assert dataframe.columns == ["name", "amount"]
    assert dataframe.row(0) == ("revenue", 100)


def test_read_excel_rejects_empty_header_with_data(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", None])
    worksheet.append(["revenue", 100])
    workbook.save(input_path)

    with pytest.raises(ValueError, match="工作表表头不能为空"):
        read_excel(input_path)


def test_read_excel_rejects_rows_with_unexpected_width(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    input_path = tmp_path / "input.xlsx"
    input_path.touch()

    class Worksheet:
        def iter_rows(self, *, values_only: bool):
            assert values_only
            return iter([("name", "amount"), ("revenue", 100, "unexpected")])

    class WorkbookStub:
        def __init__(self) -> None:
            self.sheetnames = ["Sheet1"]

        def __getitem__(self, sheet_name: str) -> Worksheet:
            assert sheet_name == "Sheet1"
            return Worksheet()

        def close(self) -> None:
            pass

    monkeypatch.setattr(
        "fin_config_calc.utils.io_dataframe_excel.load_workbook",
        lambda *args, **kwargs: WorkbookStub(),
    )

    with pytest.raises(ValueError, match=r"第 2 行有 3 列，表头有 2 列"):
        read_excel(input_path)


def test_write_excel_uses_requested_sheet_name(tmp_path: Path) -> None:
    output_path = tmp_path / "output.xlsx"

    write_excel(pl.DataFrame({"amount": [10]}), output_path, sheet_name="明细")

    workbook = load_workbook(output_path, read_only=True)
    try:
        assert workbook.sheetnames == ["明细"]
    finally:
        workbook.close()


def test_write_excel_reports_progress_by_batch(tmp_path: Path) -> None:
    output_path = tmp_path / "output.xlsx"
    progress_events: list[tuple[int, int]] = []

    write_excel(
        pl.DataFrame({"amount": [10, 20, 30, 40, 50]}),
        output_path,
        batch_size=2,
        on_progress=lambda written, total: progress_events.append((written, total)),
    )

    assert progress_events == [(2, 5), (4, 5), (5, 5)]
