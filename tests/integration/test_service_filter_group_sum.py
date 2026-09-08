from pathlib import Path

import polars as pl
from openpyxl import load_workbook

from fin_config_calc.service import filter_group_sum_to_excel


def test_filter_group_sum_to_excel_exports_comments_and_data(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    mapping_path = tmp_path / "field_mapping.xlsx"
    output_path = tmp_path / "output.xlsx"
    pl.DataFrame(
        {
            "region": ["华东", "华东", "华南"],
            "category": ["收入", "收入", "收入"],
            "amount": [100, 50, 200],
        }
    ).write_excel(input_path, worksheet="明细")
    pl.DataFrame(
        {
            "standard": ["region", "amount"],
            "comments": ["区域", "金额"],
        }
    ).write_excel(mapping_path, worksheet="字段说明")

    result = filter_group_sum_to_excel(
        input_file_path=input_path,
        input_sheet_name="明细",
        conditions={"category": "收入", "region": "华东"},
        group_sum_params={"group_by": ["region"], "sum_by": ["amount"]},
        field_mapping_file_path=mapping_path,
        field_mapping_sheet_name="字段说明",
        output_file_path=output_path,
        output_sheet_name="汇总",
    )

    assert result == output_path
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["汇总"]
        assert list(worksheet.values) == [
            ("region", "amount"),
            ("区域", "金额"),
            ("华东", 150),
        ]
        assert worksheet["B3"].data_type == "n"
    finally:
        workbook.close()


def test_filter_group_sum_to_excel_skips_grouping_when_params_are_omitted(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "input.xlsx"
    mapping_path = tmp_path / "field_mapping.xlsx"
    output_path = tmp_path / "output.xlsx"
    pl.DataFrame({"region": ["华东", "华东"], "amount": [100, 100]}).write_excel(
        input_path,
        worksheet="明细",
    )
    pl.DataFrame(
        {
            "standard": ["region", "amount"],
            "comments": ["区域", "金额"],
        }
    ).write_excel(mapping_path, worksheet="字段说明")

    filter_group_sum_to_excel(
        input_file_path=input_path,
        input_sheet_name="明细",
        conditions={},
        field_mapping_file_path=mapping_path,
        field_mapping_sheet_name="字段说明",
        output_file_path=output_path,
        output_sheet_name="明细",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert list(workbook["明细"].values) == [
            ("region", "amount"),
            ("区域", "金额"),
            ("华东", 100),
            ("华东", 100),
        ]
    finally:
        workbook.close()


def test_filter_group_sum_to_excel_skips_filtering_when_conditions_are_none(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "input.xlsx"
    mapping_path = tmp_path / "field_mapping.xlsx"
    output_path = tmp_path / "output.xlsx"
    pl.DataFrame({"region": ["华东", "华南"], "amount": [100, 200]}).write_excel(
        input_path,
        worksheet="明细",
    )
    pl.DataFrame(
        {
            "standard": ["region", "amount"],
            "comments": ["区域", "金额"],
        }
    ).write_excel(mapping_path, worksheet="字段说明")

    filter_group_sum_to_excel(
        input_file_path=input_path,
        input_sheet_name="明细",
        conditions=None,
        field_mapping_file_path=mapping_path,
        field_mapping_sheet_name="字段说明",
        output_file_path=output_path,
        output_sheet_name="明细",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert list(workbook["明细"].values) == [
            ("region", "amount"),
            ("区域", "金额"),
            ("华东", 100),
            ("华南", 200),
        ]
    finally:
        workbook.close()


def test_filter_group_sum_to_excel_omits_comment_row_without_mapping_file(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"
    pl.DataFrame({"region": ["华东"], "amount": [100]}).write_excel(
        input_path,
        worksheet="明细",
    )

    filter_group_sum_to_excel(
        input_file_path=input_path,
        input_sheet_name="明细",
        output_file_path=output_path,
        output_sheet_name="明细",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert list(workbook["明细"].values) == [
            ("region", "amount"),
            ("华东", 100),
        ]
    finally:
        workbook.close()
