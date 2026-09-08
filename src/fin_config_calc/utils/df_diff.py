import math
from collections import Counter
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import polars as pl
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import PatternFill

_DIFF_SOURCE_COLUMN = "_tmp_diff_src"


def _require_columns(dataframe: pl.DataFrame, columns: Sequence[str], table_name: str) -> None:
    missing = [column for column in columns if column not in dataframe.columns]
    if missing:
        raise ValueError(f"{table_name}缺少字段: {missing}")


def _stringify_nested(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "to_list"):
        value = value.to_list()
    return str(value)


def _normalize_for_row_diff(dataframe: pl.DataFrame) -> pl.DataFrame:
    expressions: list[pl.Expr] = []
    for column, data_type in dataframe.schema.items():
        if data_type.is_nested() or data_type == pl.Object:
            expression = pl.col(column).map_elements(
                _stringify_nested,
                return_dtype=pl.String,
                skip_nulls=False,
            )
        else:
            expression = pl.col(column).cast(pl.String, strict=False).fill_null("")
        expressions.append(expression.alias(column))
    return dataframe.with_columns(expressions)


def diff_by_row(
    df_1: pl.DataFrame,
    df_2: pl.DataFrame,
    check_columns: Sequence[str] | None = None,
    exclude_columns: Sequence[str] | None = None,
    column_map: Mapping[str, str] | None = None,
    _tmp_diff_col: str = _DIFF_SOURCE_COLUMN,
) -> pl.DataFrame:
    """按指定字段比较两个 DataFrame，返回仅出现在一侧的行。

    输入数据：
        df_1、df_2: 待比较数据；df_1 可通过 column_map 将列名对齐到 df_2。
        check_columns: 用于判断行是否相同的字段，默认使用重命名后 df_1 的全部字段。
        exclude_columns: 从比较字段中排除的字段。
        _tmp_diff_col: 输出中标识数据来源的临时字段。

    输出结果：
        包含左右两侧差异行的 DataFrame，来源字段值为 ``df_1`` 或 ``df_2``。
        空值与空字符串按相同值处理，所有字段按字符串比较，输入数据不会被修改。
    """
    if _tmp_diff_col in df_1.columns or _tmp_diff_col in df_2.columns:
        raise ValueError(f"输入数据不能包含临时来源字段: {_tmp_diff_col}")

    first = _normalize_for_row_diff(df_1.rename(dict(column_map or {})))
    second = _normalize_for_row_diff(df_2)
    compared_columns = list(check_columns) if check_columns else list(first.columns)
    if exclude_columns:
        excluded = set(exclude_columns)
        compared_columns = [column for column in compared_columns if column not in excluded]
    if not compared_columns:
        raise ValueError("用于比较的字段不能为空")
    _require_columns(first, compared_columns, "df_1")
    _require_columns(second, compared_columns, "df_2")

    combined = pl.concat(
        [
            first.with_columns(pl.lit("df_1").alias(_tmp_diff_col)),
            second.with_columns(pl.lit("df_2").alias(_tmp_diff_col)),
        ],
        how="diagonal_relaxed",
    )
    return combined.unique(subset=compared_columns, keep="none", maintain_order=True)


def _is_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and value.strip().lower() in {"nan", "none"}


def diff_by_row_pivot_to_excel(
    diff_df: pl.DataFrame,
    join_keys: Sequence[str],
    diff_source_col: str = _DIFF_SOURCE_COLUMN,
    save_path: str | Path = "diff_result.xlsx",
    sheet_name: str = "差异明细",
) -> Path:
    """将行级差异转成新旧值并列的 Excel 报表并高亮修改字段。

    输入数据：
        diff_df: ``diff_by_row`` 的输出，必须包含关联键和来源字段。
        join_keys: 将两侧记录归并到同一输出行的字段。
        diff_source_col: 标识 ``df_1``、``df_2`` 来源的字段。
        save_path、sheet_name: 输出路径模板和工作表名称。

    输出结果：
        实际生成的 Excel 路径。文件名会附加增、删、改数量；``df_1`` 显示为旧值，
        ``df_2`` 显示为新值，修改字段的旧值和新值单元格使用黄色高亮。

    使用限制：
        输出格式仅支持 ``.xlsx`` 或 ``.xlsm``。同一关联键、同一来源存在多行时，
        与参考实现一致，仅使用其中第一行。
    """
    keys = list(join_keys)
    if not keys:
        raise ValueError("join_keys 不能为空")
    _require_columns(diff_df, [*keys, diff_source_col], "diff_df")

    value_columns = [column for column in diff_df.columns if column not in {*keys, diff_source_col}]
    old_columns = [f"旧_{column}" for column in value_columns]
    new_columns = [f"新_{column}" for column in value_columns]
    output_columns = ["差异类型", *keys, *old_columns, *new_columns]
    records: list[dict[str, Any]] = []
    highlighted_columns: dict[int, set[str]] = {}
    counts: Counter[str] = Counter()

    partitions = diff_df.partition_by(keys, maintain_order=True, include_key=True, as_dict=False)
    for group in partitions:
        first_row = group.row(0, named=True)
        old_rows = group.filter(pl.col(diff_source_col) == "df_1")
        new_rows = group.filter(pl.col(diff_source_col) == "df_2")
        has_old = not old_rows.is_empty()
        has_new = not new_rows.is_empty()
        diff_type = (
            "改" if has_old and has_new else "增" if has_new else "删" if has_old else "未知"
        )
        counts[diff_type] += 1

        old_values = old_rows.row(0, named=True) if has_old else {}
        new_values = new_rows.row(0, named=True) if has_new else {}
        record = {"差异类型": diff_type, **{key: first_row[key] for key in keys}}
        changed: set[str] = set()
        for column in value_columns:
            old_value = None if _is_missing(old_values.get(column)) else old_values.get(column)
            new_value = None if _is_missing(new_values.get(column)) else new_values.get(column)
            old_column = f"旧_{column}"
            new_column = f"新_{column}"
            record[old_column] = old_value
            record[new_column] = new_value
            if diff_type == "改" and old_value != new_value:
                changed.update({old_column, new_column})
        if changed:
            highlighted_columns[len(records)] = changed
        records.append(record)

    output_path = Path(save_path)
    suffix = output_path.suffix if output_path.suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
    stem = output_path.stem or output_path.name
    final_path = output_path.with_name(
        f"{stem}(增{counts['增']}_删{counts['删']}_改{counts['改']}){suffix}"
    )
    final_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.create_sheet(title=sheet_name)
    workbook.remove(workbook["Sheet"])
    worksheet.append(output_columns)
    for record in records:
        worksheet.append([record.get(column) for column in output_columns])

    yellow_fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    column_indexes = {column: index for index, column in enumerate(output_columns, start=1)}
    for row_index, columns in highlighted_columns.items():
        for column in columns:
            worksheet.cell(row=row_index + 2, column=column_indexes[column]).fill = yellow_fill
    if records:
        worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "B2"
    workbook.save(final_path)
    logger.info("差异结果已导出: {}", final_path)
    return final_path


def diff_by_group(
    left_df: pl.DataFrame,
    left_groupby: Sequence[str],
    left_sum: Sequence[str],
    right_df: pl.DataFrame,
    right_groupby: Sequence[str] | None = None,
    right_sum: Sequence[str] | None = None,
    keep_groupby: str = "left",
    abs_delta: float = 0.5,
) -> pl.DataFrame:
    """分别分组求和后比较两个 DataFrame，返回超过差异阈值的分组。

    输入数据：
        left_df、right_df: 待汇总比较的数据。
        left_groupby、right_groupby: 两侧分组字段，字段按位置对应。
        left_sum、right_sum: 两侧求和字段，字段按位置对应。
        keep_groupby: 输出使用 ``left`` 或 ``right`` 一侧的分组字段名。
        abs_delta: 保留任一差值绝对值大于等于此阈值的分组。

    输出结果：
        分组字段、带 ``_left``/``_right`` 后缀的汇总值及带 ``_diff`` 后缀的差值。
        缺失侧在计算差值时按 0 处理，输入数据不会被修改。
    """
    left_groups = list(left_groupby)
    left_values = list(left_sum)
    right_groups = list(right_groupby) if right_groupby is not None else left_groups
    right_values = list(right_sum) if right_sum is not None else left_values
    if len(left_groups) != len(right_groups):
        raise ValueError("left_groupby 和 right_groupby 的长度不一致")
    if len(left_values) != len(right_values):
        raise ValueError("left_sum 和 right_sum 的长度不一致")
    if keep_groupby not in {"left", "right"}:
        raise ValueError("keep_groupby 必须为 'left' 或 'right'")
    if not left_groups:
        raise ValueError("groupby 字段不能为空")
    if not left_values:
        raise ValueError("sum 字段不能为空")
    _require_columns(left_df, [*left_groups, *left_values], "left_df")
    _require_columns(right_df, [*right_groups, *right_values], "right_df")

    left_aggregated = (
        left_df.select([*left_groups, *left_values])
        .with_columns(
            [
                *(pl.col(column).cast(pl.String).fill_null("None") for column in left_groups),
                *(pl.col(column).cast(pl.Float64) for column in left_values),
            ]
        )
        .group_by(left_groups, maintain_order=True)
        .agg(pl.col(column).sum().alias(f"{column}_left") for column in left_values)
    )
    right_aggregated = (
        right_df.select([*right_groups, *right_values])
        .with_columns(
            [
                *(pl.col(column).cast(pl.String).fill_null("None") for column in right_groups),
                *(pl.col(column).cast(pl.Float64) for column in right_values),
            ]
        )
        .group_by(right_groups, maintain_order=True)
        .agg(pl.col(column).sum().alias(f"{column}_right") for column in right_values)
    )

    if keep_groupby == "left":
        merge_groups = left_groups
        right_aggregated = right_aggregated.rename(
            dict(zip(right_groups, left_groups, strict=True))
        )
        diff_names = left_values
    else:
        merge_groups = right_groups
        left_aggregated = left_aggregated.rename(dict(zip(left_groups, right_groups, strict=True)))
        diff_names = right_values

    result = left_aggregated.join(
        right_aggregated,
        on=merge_groups,
        how="full",
        coalesce=True,
    )
    diff_columns: list[str] = []
    diff_expressions: list[pl.Expr] = []
    for left_column, right_column, diff_name in zip(
        left_values, right_values, diff_names, strict=True
    ):
        diff_column = f"{diff_name}_diff"
        diff_columns.append(diff_column)
        diff_expressions.append(
            (
                pl.col(f"{left_column}_left").fill_null(0)
                - pl.col(f"{right_column}_right").fill_null(0)
            ).alias(diff_column)
        )
    result = result.with_columns(diff_expressions)
    return result.filter(
        pl.any_horizontal(pl.col(column).abs() >= abs_delta for column in diff_columns)
    )
