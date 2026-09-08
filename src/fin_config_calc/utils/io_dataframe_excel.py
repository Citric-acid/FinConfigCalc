from collections.abc import Callable, Mapping
from pathlib import Path
from typing import Any, Literal, overload

import polars as pl
from loguru import logger
from openpyxl import Workbook, load_workbook


@overload
def read_excel(
    file_path: str | Path,
    sheet_name: str | None = None,
    *,
    as_dict: Literal[False] = False,
) -> pl.DataFrame: ...


@overload
def read_excel(
    file_path: str | Path,
    sheet_name: str | None = None,
    *,
    as_dict: Literal[True],
) -> dict[Any, Any]: ...


def read_excel(
    file_path: str | Path,
    sheet_name: str | None = None,
    *,
    as_dict: bool = False,
) -> pl.DataFrame | dict[Any, Any]:
    """读取 Excel 工作表内容。

    功能：
        读取指定 Excel 文件中的工作表。默认将第一行作为列名并返回 Polars
        DataFrame；当 ``as_dict=True`` 时，要求工作表恰好包含两列，并将第一列
        作为键、第二列作为值返回字典。

    输入数据：
        file_path: Excel 文件路径，文件必须存在。
        sheet_name: 要读取的工作表名称。工作簿仅有一个工作表时可以为空；存在多个
            工作表时必须指定。
        as_dict: 是否以字典形式返回，默认为 False。

    输出结果：
        ``as_dict=False`` 时返回 Polars DataFrame；``as_dict=True`` 时返回字典。

    使用限制：
        支持 ``.xlsx`` 和 ``.xlsm`` 文件。工作表第一行必须是非空且不重复的列名；
        字典模式下键不能为空或重复。
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在：{path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"不支持的 Excel 文件格式：{path.suffix}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise ValueError(
                    f"Excel 文件包含多个工作表，请指定 sheet_name：{workbook.sheetnames}"
                )
            selected_sheet_name = workbook.sheetnames[0]
        else:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"工作表不存在：{sheet_name}；可用工作表：{workbook.sheetnames}")
            selected_sheet_name = sheet_name

        rows = list(workbook[selected_sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError(f"工作表没有表头：{selected_sheet_name}")

    width = max(len(row) for row in rows)
    while width > 0 and all(
        len(row) < width or row[width - 1] is None or str(row[width - 1]).strip() == ""
        for row in rows
    ):
        width -= 1
    rows = [row[:width] for row in rows]

    headers = list(rows[0])
    if not headers or any(header is None or str(header).strip() == "" for header in headers):
        raise ValueError(f"工作表表头不能为空：{selected_sheet_name}")

    column_names = [str(header) for header in headers]
    if len(column_names) != len(set(column_names)):
        raise ValueError(f"工作表表头不能重复：{selected_sheet_name}")

    expected_width = len(column_names)
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) != expected_width:
            raise ValueError(
                f"Excel 数据列数与表头不一致：{path}，工作表 {selected_sheet_name}，"
                f"第 {row_number} 行有 {len(row)} 列，表头有 {expected_width} 列"
            )

    dataframe = pl.DataFrame(rows[1:], schema=column_names, orient="row", infer_schema_length=None)
    if not as_dict:
        return dataframe

    if dataframe.width != 2:
        raise ValueError(f"字典模式要求工作表恰好包含两列，当前为 {dataframe.width} 列")

    keys = dataframe.get_column(column_names[0]).to_list()
    if any(key is None for key in keys):
        raise ValueError("字典模式下第一列的键不能为空")
    if len(keys) != len(set(keys)):
        raise ValueError("字典模式下第一列的键不能重复")

    values = dataframe.get_column(column_names[1]).to_list()
    return dict(zip(keys, values, strict=True))


def write_excel(
    dataframe: pl.DataFrame,
    file_path: str | Path,
    sheet_name: str = "Sheet1",
    column_comments: Mapping[str, str] | None = None,
    *,
    batch_size: int = 10_000,
    on_progress: Callable[[int, int], None] | None = None,
) -> Path:
    """以低内存方式将 Polars DataFrame 写入单个 Excel 工作表。

    使用 openpyxl 的只写模式，避免 ``DataFrame.write_excel`` 在写出前调用
    ``DataFrame.rows()`` 将整张数据表物化为 Python 行元组。调用者可以通过
    ``sheet_name`` 指定工作表名称，默认名称为 ``Sheet1``。传入
    ``column_comments`` 时，英文表头下方会增加一行对应的字段说明；没有配置
    说明的字段使用原字段名。``batch_size`` 控制每批处理的行数，默认 10,000；
    每批完成后记录写入进度，并可通过 ``on_progress`` 接收“已写行数、总行数”。
    """
    path = Path(file_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("输出文件格式必须为 .xlsx")
    if batch_size <= 0:
        raise ValueError("batch_size 必须大于 0")

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(title=sheet_name)
    worksheet.append(dataframe.columns)
    if column_comments is not None:
        worksheet.append([column_comments.get(column, column) for column in dataframe.columns])
    total_rows = dataframe.height
    written_rows = 0
    for dataframe_slice in dataframe.iter_slices(n_rows=batch_size):
        for row in dataframe_slice.iter_rows(named=False, buffer_size=1):
            worksheet.append(row)
        written_rows += dataframe_slice.height
        logger.info(
            "Excel 写入进度：已写入 {}/{} 行（{:.1%}）",
            written_rows,
            total_rows,
            written_rows / total_rows,
        )
        if on_progress is not None:
            on_progress(written_rows, total_rows)
    workbook.save(path)
    return path
